#!/usr/bin/env python3
"""
Bring a farm's cattle records up to date from the farm's own spreadsheet ("Cattle list …xlsx").

    pip install openpyxl     (or: uv run --with openpyxl python3 scripts/import-cattle-xlsx.py …)
    python3 scripts/import-cattle-xlsx.py "seed_data/Cattle list Jan 2025.xlsx" > /tmp/cattle.sql 2> /tmp/cattle-report.txt
    cd apps/worker && npx wrangler d1 execute cattle-management-db --remote --file=/tmp/cattle.sql

The workbook holds dated snapshots of the herd ("cattle Feb 2025" … "cattle 09 Oct 2026") plus a
"cattle sold or died" sheet. Each animal's fields are merged oldest → newest, so the latest sheet wins.

What the SQL does for farm FARM_ID:
  - upserts every animal by ear tag (existing animal ids are kept), including animals that left before the
    first app import, with status, exit date, sale weight/price and the matching finance entry;
  - links dams by management number and calving date, and sires (Unicorn, the Simmental and Charolais bulls);
  - rebuilds calvings, services/AI, feet trimming and pelvis scoring from the sheets (these tables held only the
    earlier import of the same spreadsheet data);
  - marks Orkney Mart lots as imported where exactly one sale in the sheet matches the lot's date, weight and price.

Management numbers: cows use their own number (97, 164…), 2024 heifers keep "N.24", and calves get their
mother's number plus birth year ("97/26"; twins "102/26-A" and "-B").
"""
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta

import openpyxl

FARM_ID = 1
OWNER_USER_ID = 1

# Sheets in date order; later sheets overwrite earlier values.
SNAPSHOTS = [
    'cattle Feb 2025', 'PD calving sequence 2025', 'cattle May 2025', 'cattle Oct 2025', 'cattle Nov 2025',
    'cattle 01 Jan 2026', 'cattle size 11 Jan 2026', 'cattle 17 May 2026', 'cattle 27 jun 2026', 'cattle 09 Oct 2026',
]
LATEST = 'cattle 09 Oct 2026'
SOLD = 'cattle sold or died'

BREEDS = {'char': 'Charolais', 'ch': 'Charolais', 'char (uni)': 'Charolais', 'lim': 'Limousin', 'sim': 'Simmental',
          'simm': 'Simmental', 'aa': 'Aberdeen Angus', 'luing': 'Luing'}

# Bulls. The sheets call them by type; there was one of each on the farm.
UNICORN = 'UK523505-702252'
LEGEND = 'UK561544-100218'        # Simmental bull, listed as "UK523801-" / "Bull Sim" until his sale sheet entry
BULL_CH = 'UK523801- (Charolais bull)'  # tag never recorded
BULLS = {
    UNICORN: {'name': 'Martini Unicorn', 'label': 'Unicorn', 'breed': 'Charolais', 'origin': 'purchased'},
    LEGEND: {'name': 'Fallhills Legend', 'label': 'Legend', 'breed': 'Simmental', 'origin': 'purchased'},
    BULL_CH: {'name': 'Charolais bull', 'label': 'Bull CH', 'breed': 'Charolais', 'origin': 'purchased'},
}
SIRE_BULL = {'unicorn': UNICORN, 'sim': LEGEND, 'simm': LEGEND, 'char': BULL_CH, 'ch': BULL_CH}
SIRE_NAME = {'unicorn': 'Unicorn', 'sim': 'Simmental', 'simm': 'Simmental', 'char': 'Charolais', 'ch': 'Charolais'}

# Sales recorded outside the sheet's sale columns.
EXTRA_EXITS = {
    BULL_CH: {'status': 'sold', 'date': '2025-05-03', 'price': 3500},
}

PELVIS_DATE = '2026-04-09'
FEET_DATE = '2025-05-20'
BULL_IN_DATE_DEFAULT = '2026-04-23'

report_lines = []


def report(*a):
    report_lines.append(' '.join(str(x) for x in a))


def iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def num_label(v):
    """Management numbers come through as ints, floats (13.24, 129.1) or text ('98a')."""
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, float):
        return f'{v:.2f}'.rstrip('0').rstrip('.') if not v.is_integer() else str(int(v))
    return str(v)


def norm_tag(t):
    return re.sub(r'\s+', '', str(t)).upper()


def key(h):
    """Column header → short key."""
    h = (str(h or '')).lower().replace('\n', ' ')
    rules = [  # most specific first: "2026 calf sex" must not match "sex"
        ('2026 calf sex', 'sex26'), ('calf sex', 'sex25'), ('tag no', 'tag'), ('dam-1', 'dam1'), ('dam tag', 'dam'), ('mgmt tag', 'mgmt'), ('management tag', 'mgmt'),
        ('breed', 'breed'), ('sex', 'sex'), ('dob', 'dob'), ('size', 'size'), ('horns', 'horns'),
        ('calfed date 2024', 'calf24'), ('sire 2025 calf', 'sire25'), ('calf date 2025', 'calf25'), ('calf due 2025', 'due25'),
        ('feet trimmed', 'feet'), ('sire 2026 calf', 'sire26'),
        ('service date 2025 for 2026', 'serv25'), ('calf due date from service date', 'due26'),
        ('calf date 2026', 'calf26'), ('comment', 'comment'),
        ('ai 2025 for 2026', 'ai26'), ('bull for 2026', 'bull26'), ('service date 2026 for 2027', 'serv26'),
        ('calf due date for 2027', 'due27'), ('bull date to go', 'bullin'), ('date sold', 'exit'),
        ('sale weight', 'weight'), ('sale price', 'price'), ('yob', 'yob'),
    ]
    for needle, k in rules:
        if needle in h:
            return k
    return None


def rows(ws):
    hdr = [key(c.value) for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not clean(r[0]):
            continue
        d = {}
        for k, v in zip(hdr, r):
            if k and clean(v) is not None and k not in d:
                d[k] = clean(v)
        # "2026 calf sex" contains "calf sex": keep both apart
        yield d, hdr, r


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    animals = {}   # tag → merged dict
    seen_in = defaultdict(list)
    comments = defaultdict(list)

    for name in SNAPSHOTS:
        ws = wb[name]
        for d, hdr, raw in rows(ws):
            tag = norm_tag(d['tag'])
            if tag == 'UK523801-':
                m = str(d.get('mgmt', ''))
                if m == 'Bull Sim':
                    tag = LEGEND
                elif m == 'Bull CH':
                    tag = BULL_CH
                else:
                    continue  # calves not yet tagged when the sheet was made; they appear later with tags
            a = animals.setdefault(tag, {'tag': tag})
            seen_in[tag].append(name)
            for k, v in zip(hdr, raw):
                if k is None or clean(v) is None:
                    continue
                if k == 'exit' and not isinstance(v, datetime):
                    comments[tag].append(str(v).strip())  # the Jan 2026 sheet used this column for notes
                    continue
                if k == 'comment':
                    comments[tag].append(str(v).strip())
                    continue
                a[k] = clean(v)

    exits = {}
    for d, hdr, raw in rows(wb[SOLD]):
        tag = norm_tag(d['tag'])
        if tag == 'UK523801-':
            continue  # an untagged calf entry that is UK523801-101326 (still on the farm)
        a = animals.setdefault(tag, {'tag': tag})
        for k in ('mgmt', 'breed', 'sex', 'dob', 'yob'):
            if d.get(k) is not None and a.get(k) is None:
                a[k] = d[k]
        if d.get('sex') == 'bull':
            a['sex'] = 'bull'
        dt = iso(d.get('exit'))
        if not dt:
            continue
        w, p = d.get('weight'), d.get('price')
        exits[tag] = {'status': 'sold' if (p or w) else 'died', 'date': dt,
                      'weight': float(w) if isinstance(w, (int, float)) else None,
                      'price': float(p) if isinstance(p, (int, float)) else None,
                      'buyer': None}
        if raw[13]:
            comments[tag].append(str(raw[13]).strip())
    for tag, e in EXTRA_EXITS.items():
        exits.setdefault(tag, {'weight': None, 'price': None, 'buyer': None, **e})

    latest = {norm_tag(d['tag']) for d, _, _ in rows(wb[LATEST])}
    return animals, exits, latest, comments, seen_in


def sex_of(v):
    v = str(v or '').lower()
    if v.startswith(('fem', 'hief', 'heif')):
        return 'female', False
    if v.startswith('steer'):
        return 'male', True
    if v.startswith(('male', 'bull', 'm')):
        return 'male', False
    return None, False


def breed_of(v):
    if not v:
        return None
    return BREEDS.get(str(v).strip().lower(), str(v).strip())


def q(v):
    if v is None:
        return 'NULL'
    if isinstance(v, bool):
        return '1' if v else '0'
    if isinstance(v, (int, float)):
        return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
    return "'" + str(v).replace("'", "''") + "'"


def aid(tag):
    return f"(SELECT id FROM animals WHERE farm_id = {FARM_ID} AND tag_no = {q(tag)})"


def build(path):
    animals, exits, latest, comments, seen_in = load(path)

    # ---- basic fields
    for tag, a in animals.items():
        a['sex_v'], a['castrated'] = sex_of(a.get('sex'))
        if tag in BULLS:
            a['sex_v'], a['castrated'] = 'male', False
        a['dob_v'] = iso(a.get('dob'))
        a['yob_v'] = int(a['yob']) if isinstance(a.get('yob'), (int, float)) else (int(a['dob_v'][:4]) if a['dob_v'] else None)
        a['breed_v'] = breed_of(a.get('breed')) or BULLS.get(tag, {}).get('breed')
        a['mgmt_raw'] = num_label(a.get('mgmt'))
        a['dam_no'] = num_label(a.get('dam'))
        a['dam1_no'] = num_label(a.get('dam1'))
        a['uni'] = '(uni)' in str(a.get('breed', '')).lower()
        if any(str(x).lower().startswith('steer') for x in [a.get('sex')]):
            a['castrated'] = True

    # Animals listed as steers in any sheet were castrated
    wb = openpyxl.load_workbook(path, data_only=True)
    for name in ['steer growth rates 2025', 'cattle Oct 2025', 'cattle Nov 2025', 'cattle May 2025', 'cattle Feb 2025']:
        for d, _, _ in rows(wb[name]):
            if str(d.get('sex', '')).lower().startswith('steer') and norm_tag(d['tag']) in animals:
                animals[norm_tag(d['tag'])]['castrated'] = True

    # Every crop of home-bred males has been sold as steers, so treat home-bred males as castrated.
    for tag, a in animals.items():
        if a['sex_v'] == 'male' and tag not in BULLS and tag.startswith('UK523801-'):
            a['castrated'] = True

    # ---- who owns which management number: breeding females (and the 2024 heifers' N.24 numbers)
    def owns_number(a):
        m = a['mgmt_raw']
        if not m or a['sex_v'] != 'female' or a['tag'] in BULLS:
            return False
        if re.fullmatch(r'\d+\.24', m):
            return True
        if a['yob_v'] and a['yob_v'] <= 2022:
            return True
        # 2023 heifers kept for breeding have their own number in the 2026 sheets
        in_2026 = any('2026' in n for n in seen_in[a['tag']])
        return in_2026 and m != a['dam_no'] and a['yob_v'] and a['yob_v'] <= 2023

    owner = {}
    for a in sorted(animals.values(), key=lambda x: x['dob_v'] or '9999'):
        if owns_number(a):
            if a['mgmt_raw'] in owner:
                report('number used twice:', a['mgmt_raw'], owner[a['mgmt_raw']], a['tag'])
            else:
                owner[a['mgmt_raw']] = a['tag']
            a['own'] = a['mgmt_raw']

    # ---- calvings recorded on each cow (date, sex text, sire text) per year
    calvings = defaultdict(list)  # cow tag → [(date, sexText, sireText, year)]
    for tag, a in animals.items():
        for dk, sk, rk, yr in (('calf24', None, None, 2024), ('calf25', 'sex25', 'sire25', 2025), ('calf26', 'sex26', 'sire26', 2026)):
            d = iso(a.get(dk))
            if not d and yr == 2025:
                d = iso(a.get('due25')) if 'cattle Oct 2025' in seen_in[tag] else None
            if d:
                calvings[tag].append({'date': d, 'sex': a.get(sk) if sk else None, 'sire': a.get(rk) if rk else None, 'year': yr})

    # ---- link calves to dams: the cow whose number the calf carries and whose calving date matches its DOB
    for tag, a in animals.items():
        if not a['dob_v'] or a.get('own') and a['yob_v'] and a['yob_v'] <= 2022 and not a['dam_no']:
            pass
        cands = [n for n in (a['mgmt_raw'], a['dam_no']) if n and n != a.get('own')]
        dam = None
        if a['dob_v']:
            dob = date.fromisoformat(a['dob_v'])
            for n in cands:
                cow = owner.get(n)
                if cow and cow != tag and any(abs((date.fromisoformat(c['date']) - dob).days) <= 7 for c in calvings[cow]):
                    dam = cow
                    break
        if not dam and a['dam_no'] and owner.get(a['dam_no']) not in (None, tag):
            cow = animals[owner[a['dam_no']]]
            if cow['dob_v'] and a['dob_v'] and cow['dob_v'] < a['dob_v']:
                dam = owner[a['dam_no']]
        a['dam_tag'] = dam

    # ---- sires for calves from the dam's record for that year
    for tag, a in animals.items():
        a['sire_tag'], a['sire_name'] = None, None
        if a['uni']:
            a['sire_tag'], a['sire_name'] = UNICORN, 'Unicorn'
        dam = animals.get(a['dam_tag']) if a['dam_tag'] else None
        if dam and a['dob_v']:
            for c in calvings[dam['tag']]:
                if abs((date.fromisoformat(c['date']) - date.fromisoformat(a['dob_v'])).days) <= 7 and c['sire']:
                    s = str(c['sire']).strip().lower()
                    a['sire_tag'] = a['sire_tag'] or SIRE_BULL.get(s)
                    a['sire_name'] = a['sire_name'] or SIRE_NAME.get(s, str(c['sire']).strip())
        if a['sire_tag'] and a['dob_v'] and animals.get(a['sire_tag'], {}).get('dob_v', '0') >= a['dob_v']:
            a['sire_tag'] = None

    # ---- management labels
    by_dam_year = defaultdict(list)
    for tag, a in animals.items():
        if tag in BULLS:
            a['label'] = BULLS[tag]['label']
            continue
        if a.get('own'):
            a['label'] = a['own']
            continue
        dam_no = animals[a['dam_tag']]['own'] if a['dam_tag'] else (a['mgmt_raw'] if a['mgmt_raw'] and re.fullmatch(r'\d+', a['mgmt_raw']) else a['dam_no'])
        if a['mgmt_raw'] and re.fullmatch(r'\d+[ab]', a['mgmt_raw']):  # 2024 twins '98a'/'98b'
            dam_no = a['mgmt_raw'][:-1]
        if dam_no and a['yob_v']:
            a['label'] = f"{dam_no}/{str(a['yob_v'])[2:]}"
            by_dam_year[a['label']].append(a)
        else:
            a['label'] = a['mgmt_raw']
    for label, group in by_dam_year.items():
        if len(group) > 1:
            for i, a in enumerate(sorted(group, key=lambda x: (x['dob_v'] or '', x['tag']))):
                a['label'] = f'{label}-{"ABCDEFGH"[i]}'
            report('twins/same label:', label, [g['tag'] for g in group])

    # ---- notes
    for tag, a in animals.items():
        notes = []
        for c in dict.fromkeys(comments.get(tag, [])):
            notes.append(c)
        if a.get('horns') in ('horned', 'polled'):
            notes.append(a['horns'].capitalize())
        if not a['dam_tag'] and a['dam_no'] and tag not in BULLS:
            notes.append(f"Dam no. {a['dam_no']}" + (f", grand-dam no. {a['dam1_no']}" if a['dam1_no'] else '') + ' (not on record)')
        a['notes'] = '\n'.join(notes) or None

    return animals, exits, latest, calvings


def sql(path, existing_tags, mart_dates=()):
    animals, exits, latest, calvings = build(path)
    # No price on the sheet but the farm sold at the mart that day: sold, price still to match from the mart report
    for e in exits.values():
        if e['status'] == 'died' and e['date'] in set(mart_dates):
            e['status'], e['unpriced'] = 'sold', True
    out = ['-- Generated by scripts/import-cattle-xlsx.py', 'PRAGMA defer_foreign_keys = true;']

    # The Simmental bull was imported earlier with the incomplete tag "UK523801-".
    if 'UK523801-' in existing_tags:
        out.append(f"UPDATE animals SET tag_no = {q(LEGEND)} WHERE farm_id = {FARM_ID} AND tag_no = 'UK523801-';")

    # ---- animals
    for tag, a in sorted(animals.items(), key=lambda kv: kv[1]['dob_v'] or ''):
        e = exits.get(tag)
        if tag not in latest and not e:
            report('not in latest sheet and no exit recorded — kept active:', tag, a.get('label'))
        if tag in latest and e:
            report('in latest sheet but also on sold sheet — treated as sold:', tag, e)
        status = e['status'] if e else 'active'
        bull = BULLS.get(tag, {})
        row = {
            'farm_id': FARM_ID, 'species': 'cattle', 'tag_no': tag, 'management_tag': a.get('label'),
            'name': bull.get('name'), 'sex': a['sex_v'] or 'female', 'castrated': a['castrated'], 'breed': a['breed_v'],
            'dob': a['dob_v'], 'birth_year': a['yob_v'], 'size': int(a['size']) if isinstance(a.get('size'), (int, float)) else None,
            'sire_name': a['sire_name'], 'status': status, 'origin': bull.get('origin', 'born' if tag.startswith('UK523801') else 'purchased'),
            'exit_date': e['date'] if e else None,
            'exit_reason': (('Sold at Orkney Mart; price not recorded' if e.get('unpriced') else 'Sold') if status == 'sold' else 'Died') if e else None,
            'sale_price': e['price'] if e else None, 'sale_weight_kg': e['weight'] if e else None,
            'notes': a['notes'],
        }
        cols = ', '.join(row)
        vals = ', '.join(q(v) for v in row.values())
        upd = ', '.join(f'{k} = excluded.{k}' for k in row if k not in ('farm_id', 'tag_no')) + ", updated_at = datetime('now')"
        out.append(f'INSERT INTO animals ({cols}) VALUES ({vals}) ON CONFLICT (farm_id, tag_no) DO UPDATE SET {upd};')

    # ---- parents
    for tag, a in animals.items():
        out.append(f"UPDATE animals SET dam_id = {aid(a['dam_tag']) if a['dam_tag'] else 'NULL'}, "
                   f"sire_id = {aid(a['sire_tag']) if a['sire_tag'] else 'NULL'} WHERE farm_id = {FARM_ID} AND tag_no = {q(tag)};")

    # ---- records rebuilt from the sheets
    for t in ('birth_records', 'breeding_records', 'health_records', 'weight_records'):
        out.append(f'DELETE FROM {t} WHERE farm_id = {FARM_ID};')
    out.append(f"DELETE FROM transactions WHERE farm_id = {FARM_ID} AND source = 'animal_sale';")

    # Calvings: one birth record per calf (twins → two), linked to the calf where we have it.
    calves_of = defaultdict(list)
    for tag, a in animals.items():
        if a['dam_tag']:
            calves_of[a['dam_tag']].append(a)
    births = 0
    for cow, cs in calvings.items():
        for c in cs:
            d = date.fromisoformat(c['date'])
            kids = [k for k in calves_of[cow] if k['dob_v'] and abs((date.fromisoformat(k['dob_v']) - d).days) <= 7]
            text = str(c['sex'] or '').lower()
            outcome = 'stillborn' if 'doa' in text else 'alive'
            notes = c['sex'] if c['sex'] and not re.fullmatch(r'(fe)?male|m|f', text) else None
            sire_s = str(c['sire'] or '').strip().lower()
            sire_tag = SIRE_BULL.get(sire_s)
            sire_name = SIRE_NAME.get(sire_s, c['sire'])
            n = 2 if ('twin' in text or '+' in text) else 1
            sexes = []
            if '+' in text:
                sexes = ['male' if p.strip().startswith('m') else 'female' for p in text.split('+')]
            else:
                sexes = ['female' if text.startswith('f') else 'male' if text.startswith('m') else None] * n
            kids = sorted(kids, key=lambda k: k['tag'])
            for i in range(max(n, len(kids))):
                kid = kids[i] if i < len(kids) else None
                sx = kid['sex_v'] if kid else (sexes[i] if i < len(sexes) else None)
                row = {'farm_id': FARM_ID, 'dam_id': aid(cow), 'offspring_id': aid(kid['tag']) if kid else None, 'date': c['date'],
                       'sex': sx, 'sire_id': aid(sire_tag) if sire_tag and (kid is None or kid['sire_tag']) else None,
                       'sire_name': (kid['sire_name'] if kid and kid['sire_name'] else sire_name), 'outcome': outcome, 'notes': notes,
                       'created_by': OWNER_USER_ID}
                vals = ', '.join(v if k in ('dam_id', 'offspring_id', 'sire_id') and v else q(v) for k, v in row.items())
                out.append(f"INSERT INTO birth_records ({', '.join(row)}) VALUES ({vals});")
                births += 1
    report('birth records:', births)

    # Services. 2025 (for 2026 calves): natural, sire from "Sire 2026 calf". 2026 (for 2027): AI and/or bull.
    services = 0
    for tag, a in animals.items():
        s25 = iso(a.get('serv25'))
        if s25:
            sire = str(a.get('sire26') or '').strip().lower()
            calved = any(c['year'] == 2026 for c in calvings[tag])
            row = {'animal_id': aid(tag), 'event': 'service', 'date': s25, 'method': 'natural', 'sire_id': aid(SIRE_BULL[sire]) if sire in SIRE_BULL else None,
                   'sire_name': SIRE_NAME.get(sire, a.get('sire26')), 'expected_due_date': iso(a.get('due26')),
                   'outcome': 'born' if calved else 'pending', 'notes': None}
            out.append(service_sql(row)); services += 1
        ai = a.get('ai26')
        s26 = iso(a.get('serv26'))
        bull = str(a.get('bull26') or '').strip()
        bull_in = iso(a.get('bullin')) or (BULL_IN_DATE_DEFAULT if bull.lower() in ('unicorn', 'simm') else None)
        bull_tag = UNICORN if bull.lower() == 'unicorn' else LEGEND if bull.lower() == 'simm' else None
        bull_name = 'Unicorn' if bull_tag == UNICORN else 'Fallhills Legend (Simmental)' if bull_tag else None
        if ai and str(ai).lower() != 'no' and s26:
            row = {'animal_id': aid(tag), 'event': 'service', 'date': s26, 'method': 'ai', 'sire_id': None,
                   'sire_name': str(ai).strip(), 'expected_due_date': iso(a.get('due27')), 'outcome': 'pending',
                   'notes': f'{bull_name} ran with the group from {bull_in} as a sweeper' if bull_tag else None}
            out.append(service_sql(row)); services += 1
        elif bull_tag and a['tag'] not in BULLS and tag not in exits and a['sex_v'] == 'female':
            due = (date.fromisoformat(bull_in) + timedelta(days=283)).isoformat()
            row = {'animal_id': aid(tag), 'event': 'service', 'date': bull_in, 'method': 'natural', 'sire_id': aid(bull_tag),
                   'sire_name': bull_name, 'expected_due_date': due, 'outcome': 'pending',
                   'notes': 'Bull put in with the group; due date assumes service on the first day'}
            out.append(service_sql(row)); services += 1
    report('service records:', services)

    # Health: feet trimming and pelvis scoring
    health = 0
    for tag, a in animals.items():
        if str(a.get('feet') or '').lower().startswith('y'):
            out.append(health_sql(tag, FEET_DATE, 'procedure', 'Feet trimmed', None)); health += 1
        for c in a.get('notes', '') .split('\n') if a.get('notes') else []:
            m = re.match(r'pelvis measure (\w+)', c)
            if m:
                out.append(health_sql(tag, PELVIS_DATE, 'test', 'Pelvic measurement', m.group(1))); health += 1
            if c.lower().startswith('johnes positive'):
                out.append(health_sql(tag, '2025-01-15', 'test', "Johne's test", 'Positive, then confirmed negative', notes='Date approximate (January 2025)')); health += 1
    report('health records:', health)

    # Sales: weight record + finance entry (same shape as services/finance.ts)
    for tag, e in exits.items():
        if e['weight']:
            out.append(f"INSERT INTO weight_records (farm_id, animal_id, date, weight_kg, notes, created_by) VALUES ({FARM_ID}, {aid(tag)}, {q(e['date'])}, {q(e['weight'])}, 'Sale weight', {OWNER_USER_ID});")
        if e['status'] == 'sold' and e['price']:
            out.append("INSERT INTO transactions (farm_id, type, category, amount, date, description, animal_id, source, source_id, created_by) "
                       f"VALUES ({FARM_ID}, 'income', 'livestock_sales', {q(e['price'])}, {q(e['date'])}, {q('Sale of ' + tag)}, {aid(tag)}, 'animal_sale', {aid(tag)}, {OWNER_USER_ID});")

    return out, animals, exits


def service_sql(row):
    full = {'farm_id': FARM_ID, **row, 'created_by': OWNER_USER_ID}
    vals = ', '.join(v if k in ('animal_id', 'sire_id') and v else q(v) for k, v in full.items())
    return f"INSERT INTO breeding_records ({', '.join(full)}) VALUES ({vals});"


def health_sql(tag, d, typ, title, result, notes=None):
    return ("INSERT INTO health_records (farm_id, animal_id, date, type, title, result, notes, created_by) "
            f"VALUES ({FARM_ID}, {aid(tag)}, {q(d)}, {q(typ)}, {q(title)}, {q(result)}, {q(notes)}, {OWNER_USER_ID});")


def mart_links(animals, exits, lots):
    """Lots whose date, weight and price match exactly one sale in the sheet."""
    out = []
    for lot in lots:
        hits = [t for t, e in exits.items() if e['date'] == lot['date'] and e['price'] == lot['price']
                and (lot['weight_kg'] is None or e['weight'] == lot['weight_kg'])]
        if len(hits) == 1:
            tag = hits[0]
            out.append(f"UPDATE mart_matches SET status = 'imported', animal_ids = json_array({aid(tag)}), decided_by = {OWNER_USER_ID}, "
                       f"decided_at = datetime('now') WHERE id = {lot['mid']};")
            out.append(f"UPDATE animals SET buyer = COALESCE(buyer, 'Orkney Auction Mart') WHERE farm_id = {FARM_ID} AND tag_no = {q(tag)};")
            out.append(f"UPDATE transactions SET counterparty = 'Orkney Auction Mart' WHERE farm_id = {FARM_ID} AND source = 'animal_sale' AND source_id = {aid(tag)};")
            report('mart lot', lot['mid'], lot['date'], lot['weight_kg'], lot['price'], '→', tag, animals[tag].get('label'))
        else:
            report('mart lot', lot['mid'], lot['date'], lot['weight_kg'], lot['price'], 'left to review:', len(hits), 'candidates')
    return out


if __name__ == '__main__':
    path = sys.argv[1]
    existing = set(json.load(open(sys.argv[2]))) if len(sys.argv) > 2 else set()
    lots = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else []
    out, animals, exits = sql(path, existing, {lot['date'] for lot in lots})
    out += mart_links(animals, exits, lots)
    print('\n'.join(out))
    active = sum(1 for t in animals if t not in exits)
    report(f'animals: {len(animals)} ({active} on farm, {len(exits)} sold or died)')
    print('\n'.join(report_lines), file=sys.stderr)
